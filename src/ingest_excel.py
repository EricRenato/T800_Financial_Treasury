"""Read contract rows from the local workbook without changing the source file."""

import argparse
from pathlib import Path

import openpyxl

from common import CONTRACTS_PATH
from contract_engine import create_and_unfold_contract, validate_contract


def load_contracts_from_excel(excel_path=CONTRACTS_PATH):
    excel_path = Path(excel_path)
    if not excel_path.exists():
        raise FileNotFoundError(f"planilha de contratos nao encontrada: {excel_path}")

    workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if "contracts" not in workbook.sheetnames:
            raise ValueError("a planilha precisa conter uma aba chamada 'contracts'")
        sheet = workbook["contracts"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError("a aba 'contracts' esta vazia")
        columns = {str(header).strip().lower(): index for index, header in enumerate(headers) if header}
        required = {"contract_id", "bank_id", "start_date", "settlement_date", "due_date", "principal", "rate_cdi"}
        missing = sorted(required - columns.keys())
        if missing:
            raise ValueError("colunas obrigatorias ausentes: " + ", ".join(missing))

        contracts = []
        for row_number, row in enumerate(rows, start=2):
            if not row or all(value is None or str(value).strip() == "" for value in row):
                continue
            payload = {name: row[index] if index < len(row) else None for name, index in columns.items()}
            if not str(payload.get("asset_type") or "").strip():
                payload["asset_type"] = "CDB"
            try:
                contracts.append((row_number, validate_contract(payload)))
            except (TypeError, ValueError) as exc:
                raise ValueError(f"linha {row_number}: {exc}") from exc
        return contracts
    finally:
        workbook.close()


def ingest_contracts_from_excel(excel_path=CONTRACTS_PATH, db_path=None, contract_id=None):
    contracts = load_contracts_from_excel(excel_path)
    if contract_id:
        contracts = [item for item in contracts if item[1]["contract_id"] == contract_id]
        if not contracts:
            raise ValueError(f"contrato {contract_id!r} nao encontrado na planilha")

    summaries = []
    for row_number, contract in contracts:
        result = create_and_unfold_contract(contract, db_path) if db_path else create_and_unfold_contract(contract)
        result["source_row"] = row_number
        summaries.append(result)
    return summaries


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Processa contratos da planilha local")
    parser.add_argument("--input", type=Path, default=CONTRACTS_PATH, help="caminho do Excel de contratos")
    parser.add_argument("--contract-id", help="processa apenas este ID")
    args = parser.parse_args()
    for item in ingest_contracts_from_excel(args.input, contract_id=args.contract_id):
        print(f"[OK] {item['contract_id']}: {item['installment_count']} parcelas, "
              f"{item['accounting_count']} lancamentos contabeis (linha {item['source_row']})")
